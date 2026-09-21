"""File loaders. Each turns one file into a single Markdown-flavoured `Document`.

Everything is normalised to lightweight Markdown (`#` headings, `|` tables, `-` lists) so the
structure-aware chunker treats every format the same way. PDF pages are separated by a form feed
(`\\f`), which the chunkers use to record page numbers.
"""

from __future__ import annotations

import asyncio
import csv
import re
from collections.abc import AsyncIterator
from pathlib import Path
from typing import Any

from kbsdk.adapters._deps import require
from kbsdk.errors import ConfigError
from kbsdk.interfaces import OCREngine
from kbsdk.text import stable_hash
from kbsdk.types import Document

EXTENSION_LOADERS: dict[str, str] = {
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
    ".rst": "text",
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".csv": "csv",
    ".html": "html",
    ".htm": "html",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".tif": "image",
    ".tiff": "image",
    ".bmp": "image",
    ".webp": "image",
}


def _import(module: str, extra: str, provider: str) -> Any:
    return require(module, stage="loader", provider=provider, extra=extra)


def to_markdown_table(rows: list[list[str]]) -> str:
    rows = [
        [" ".join(cell.split()).replace("|", "\\|") for cell in row] for row in rows if any(row)
    ]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(padded[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(row) + " |" for row in padded[1:]]
    return "\n".join(lines)


class _FileLoader:
    """Base: subclasses implement `_parse(path) -> (markdown, metadata)`.

    `extraction` and `ocr` come from the knowledge config (see `KnowledgeConfig`); loaders that have
    no use for them simply ignore them.
    """

    file_type = ""

    def __init__(self, *, extraction: str = "layout", ocr: OCREngine | None = None) -> None:
        self.extraction = extraction
        self.ocr = ocr

    def _parse(self, path: Path) -> tuple[str, dict[str, Any]]:
        raise NotImplementedError

    def _document(self, path: Path, text: str, meta: dict[str, Any]) -> Document:
        title = str(meta.pop("title", "") or path.stem)
        return Document(
            id=stable_hash(str(path.resolve()), length=16),
            source=path.name,
            text=text,
            metadata={"title": title, "file_type": self.file_type, **meta},
        )

    async def load(self, location: str) -> AsyncIterator[Document]:
        path = Path(location)
        text, meta = await asyncio.to_thread(self._parse, path)
        yield self._document(path, text, meta)


class TextLoader(_FileLoader):
    file_type = "text"

    def _parse(self, path: Path) -> tuple[str, dict[str, Any]]:
        text = path.read_text(encoding="utf-8", errors="replace")
        heading = re.search(r"^#\s+(.+?)\s*$", text, re.MULTILINE)
        return text, {"title": heading.group(1) if heading else ""}


_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)+)[.)]?\s+[A-Z].{2,80}$")
_SECTION_HEADING = re.compile(
    r"^(?:section|article|chapter|part)\s+\d+[.:]?\s+.{2,80}$", re.IGNORECASE
)
_ALL_CAPS_HEADING = re.compile(r"^[A-Z][A-Z0-9 &/,\-']{3,59}$")


def promote_headings(text: str) -> str:
    """Turn heading-looking lines of plain text (PDF) into Markdown headings.

    Deliberately conservative: only dotted numbering ("3.2 Casual leave"), "Section 4 ..." lines and
    short ALL-CAPS lines qualify, so numbered list items ("1. Submit the form") are left alone.
    """
    # Pages are separated by a form feed, which would otherwise be glued to the first line of the
    # next page and hide a heading there. So each page is processed on its own.
    return "\f".join(_promote_page(page) for page in text.split("\f"))


def _promote_page(text: str) -> str:
    out: list[str] = []
    for line in text.split("\n"):
        stripped = line.strip()
        level = 0
        if stripped and len(stripped) <= 90 and not stripped.endswith((".", ":", ",")):
            if numbered := _NUMBERED_HEADING.match(stripped):
                level = 1 + numbered.group(1).count(".")
            elif _SECTION_HEADING.match(stripped) or (
                _ALL_CAPS_HEADING.match(stripped) and len(stripped.split()) >= 2
            ):
                level = 1
        out.append(f"{'#' * level} {stripped}" if level else line)
    return "\n".join(out)


MIN_PAGE_CHARS = 20  # a page with less text than this is treated as scanned
RENDER_DPI = 250  # resolution pages are rendered at before OCR


class PdfLoader(_FileLoader):
    """PDFs.

    `extraction="layout"` (the default) reads tables as Markdown tables and keeps reading order;
    `"text"` is faster but flattens tables into running text. Pages with no text layer are read with
    OCR when an OCR engine is configured; otherwise they are flagged (`needs_ocr`).
    """

    file_type = "pdf"

    def _text_pages(self, path: Path) -> tuple[list[str], dict[str, Any]]:
        pypdf = _import("pypdf", "pdf", "pdf")
        reader = pypdf.PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF is password-protected")
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        return pages, {"title": reader.metadata.title if reader.metadata else ""}

    def _layout_pages(self, path: Path) -> tuple[list[str], dict[str, Any]]:
        pdfplumber = _import("pdfplumber", "pdf", "pdf")
        try:
            pdf = pdfplumber.open(str(path))
        except Exception as exc:
            if "password" in type(exc).__name__.lower():
                raise ValueError("PDF is password-protected") from exc
            raise
        with pdf:
            pages = [self._layout_page(page) for page in pdf.pages]
            title = str((pdf.metadata or {}).get("Title") or "")
        return pages, {"title": title}

    @staticmethod
    def _layout_page(page: Any) -> str:
        tables = page.find_tables()
        boxes = [t.bbox for t in tables]

        def outside_tables(obj: dict[str, Any]) -> bool:
            return not any(
                obj["x0"] >= x0 - 1
                and obj["x1"] <= x1 + 1
                and obj["top"] >= top - 1
                and obj["bottom"] <= bottom + 1
                for x0, top, x1, bottom in boxes
            )

        body = page.filter(outside_tables) if boxes else page
        items: list[tuple[float, str]] = [
            (float(line["top"]), str(line["text"])) for line in body.extract_text_lines()
        ]
        for table in tables:
            rows = [[cell or "" for cell in row] for row in table.extract()]
            markdown = to_markdown_table(rows)
            if markdown:
                items.append((float(table.bbox[1]), f"\n{markdown}\n"))
        items.sort(key=lambda item: item[0])
        return "\n".join(text for _, text in items).strip()

    def _extract(self, path: Path) -> tuple[list[str], dict[str, Any]]:
        if self.extraction == "text":
            return self._text_pages(path)
        return self._layout_pages(path)

    def _render(self, path: Path, indexes: list[int]) -> dict[int, Any]:
        pdfium = _import("pypdfium2", "pdf", "pdf")
        pdf = pdfium.PdfDocument(str(path))
        try:
            return {i: pdf[i].render(scale=RENDER_DPI / 72).to_pil() for i in indexes}
        finally:
            pdf.close()

    async def load(self, location: str) -> AsyncIterator[Document]:
        path = Path(location)
        pages, meta = await asyncio.to_thread(self._extract, path)
        meta["pages"] = len(pages)
        blank = [i for i, text in enumerate(pages) if len(text.strip()) < MIN_PAGE_CHARS]
        if blank and self.ocr is not None:
            images = await asyncio.to_thread(self._render, path, blank)
            recovered = 0
            for index in blank:
                text = (await self.ocr.recognize(images[index])).strip()
                if len(text) > len(pages[index].strip()):  # never replace real text with less
                    pages[index] = text
                    recovered += 1
            meta["ocr_pages"] = recovered
            blank = [i for i in blank if not pages[i].strip()]
        if blank:
            meta["needs_ocr"] = True
            meta["pages_without_text"] = len(blank)
        text = promote_headings("\f".join(pages))
        yield self._document(path, text, meta)


class ImageLoader(_FileLoader):
    """Photos and scans of documents (PNG, JPEG, TIFF, ...). Needs an OCR engine (`knowledge.ocr`)."""

    file_type = "image"

    async def load(self, location: str) -> AsyncIterator[Document]:
        path = Path(location)
        if self.ocr is None:
            raise ConfigError(
                f"{path.name}: reading images needs OCR. Set `knowledge.ocr` "
                "(e.g. {provider: rapidocr}) and pip install 'kbsdk[ocr]'."
            )
        frames = await asyncio.to_thread(self._frames, path)
        texts = [(await self.ocr.recognize(frame)).strip() for frame in frames]
        meta: dict[str, Any] = {"pages": len(frames), "ocr_pages": sum(1 for t in texts if t)}
        yield self._document(path, promote_headings("\f".join(texts)), meta)

    @staticmethod
    def _frames(path: Path) -> list[Any]:
        image_module = _import("PIL.Image", "ocr", "image")
        frames: list[Any] = []
        with image_module.open(path) as image:
            for index in range(getattr(image, "n_frames", 1)):  # multi-page TIFFs
                image.seek(index)
                frames.append(image.convert("RGB"))
        return frames


class DocxLoader(_FileLoader):
    file_type = "docx"

    def _parse(self, path: Path) -> tuple[str, dict[str, Any]]:
        docx = _import("docx", "office", "docx")
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(str(path))
        blocks: list[str] = []
        title = document.core_properties.title or ""
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if not text:
                    continue
                style = (item.style.name if item.style is not None else "") or ""
                if style == "Title":
                    title = title or text
                    blocks.append(f"# {text}")
                elif style.startswith("Heading") and style.split()[-1].isdigit():
                    blocks.append(f"{'#' * min(6, int(style.split()[-1]))} {text}")
                elif style.startswith("List"):
                    blocks.append(f"- {text}")
                else:
                    blocks.append(text)
            elif isinstance(item, Table):
                rows = [[cell.text for cell in row.cells] for row in item.rows]
                blocks.append(to_markdown_table(rows))
        return "\n\n".join(b for b in blocks if b), {"title": title}


class PptxLoader(_FileLoader):
    file_type = "pptx"

    def _parse(self, path: Path) -> tuple[str, dict[str, Any]]:
        pptx = _import("pptx", "office", "pptx")
        presentation = pptx.Presentation(str(path))
        blocks: list[str] = []
        for number, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            title = title_shape.text_frame.text.strip() if title_shape is not None else ""
            blocks.append(f"## Slide {number}: {title}" if title else f"## Slide {number}")
            for shape in slide.shapes:
                if shape == title_shape:
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    blocks.append(to_markdown_table(rows))
                elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    text = "\n".join(
                        p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
                    )
                    if text:
                        blocks.append(text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    blocks.append(f"Speaker notes: {notes}")
        return "\n\n".join(blocks), {
            "title": presentation.core_properties.title or "",
            "slides": len(presentation.slides),
        }


class XlsxLoader(_FileLoader):
    file_type = "xlsx"

    def _parse(self, path: Path) -> tuple[str, dict[str, Any]]:
        openpyxl = _import("openpyxl", "office", "xlsx")
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        blocks: list[str] = []
        try:
            for sheet in workbook.worksheets:
                rows = [
                    ["" if v is None else str(v) for v in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                table = to_markdown_table(rows)
                if table:
                    blocks.append(f"## Sheet: {sheet.title}\n\n{table}")
            sheets = len(workbook.worksheets)
        finally:
            workbook.close()
        return "\n\n".join(blocks), {"sheets": sheets}


class CsvLoader(_FileLoader):
    file_type = "csv"

    def _parse(self, path: Path) -> tuple[str, dict[str, Any]]:
        raw = path.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
                raw[:4096], delimiters=",;\t|"
            )
        except csv.Error:
            dialect = csv.excel
        rows = [list(row) for row in csv.reader(raw.splitlines(), dialect)]
        return f"## {path.stem}\n\n{to_markdown_table(rows)}", {}


_DROP_TAGS = ("script", "style", "noscript", "nav", "footer", "aside", "form", "svg", "template")
_HEADINGS = {f"h{n}": n for n in range(1, 7)}
_TEXT_BLOCKS = {"p", "pre", "blockquote", "dt", "dd", "summary", "figcaption"}


class HtmlLoader(_FileLoader):
    """HTML pages and wiki exports. Navigation, scripts and footers are stripped."""

    file_type = "html"

    def _parse(self, path: Path) -> tuple[str, dict[str, Any]]:
        bs4 = _import("bs4", "web", "html")
        soup = bs4.BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
        title = soup.title.get_text(strip=True) if soup.title else ""
        for tag in soup(_DROP_TAGS):
            tag.decompose()
        blocks: list[str] = []
        self._walk(soup.body or soup, blocks)
        if not title:
            first = next((b for b in blocks if b.startswith("# ")), "")
            title = first[2:]
        return "\n\n".join(b for b in blocks if b), {"title": title}

    def _walk(self, node: Any, blocks: list[str]) -> None:
        for child in node.children:
            name = getattr(child, "name", None)
            if name is None:
                continue
            if name in _HEADINGS:
                text = child.get_text(" ", strip=True)
                if text:
                    blocks.append(f"{'#' * _HEADINGS[name]} {text}")
            elif name in _TEXT_BLOCKS:
                text = child.get_text(" ", strip=True)
                if text:
                    blocks.append(text)
            elif name == "li":
                own = " ".join(
                    part.get_text(" ", strip=True)
                    if hasattr(part, "get_text")
                    else str(part).strip()
                    for part in child.children
                    if getattr(part, "name", None) not in {"ul", "ol"}
                ).strip()
                if own:
                    blocks.append(f"- {own}")
                self._walk(child, blocks)
            elif name == "table":
                rows = [
                    [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
                    for row in child.find_all("tr")
                ]
                blocks.append(to_markdown_table(rows))
            else:
                self._walk(child, blocks)
